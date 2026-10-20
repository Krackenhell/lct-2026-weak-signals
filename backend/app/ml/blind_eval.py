"""Слепой замер на новых отраслях (раздел 7 исследования): одна честная цифра «доля полезных карточек в топ-15».

1. export — лист Excel для разметчиков: карточки топ-15 нескольких прогонов (наш конвейер и, если есть, «только LLM»)
   перемешаны, без оценки модели, без места и без происхождения. Рубрика по каждой карточке:
   релевантность 0–2, технология 0–2 (0 — заголовок статьи, организация, продукт), ранняя стадия 0–2 (0 — зрелое или стандарт),
   русское название 0–2, категория, комментарий. Ключ (какая карточка откуда) — в отдельном файле, разметчику не показывается.
2. score — читает заполненные листы и считает:
   «полезная» карточка = релевантность ≥ 1, технология = 2, ранняя стадия ≥ 1;
   U@15 — число полезных в топ-15 на запрос, P@15 — доля полезных при фиксированном знаменателе 15;
   по вариантам (конвейер / только LLM) и по запросам; при двух разметчиках — согласие (каппа Коэна).

Запуск:
    python -m app.ml.blind_eval run-holdout [--with-llm] [--only "q1;q2"]   # живые прогоны тем HOLD-OUT (data/bench/queries.json)
    python -m app.ml.blind_eval export --runs 61,62,63    # или --queries "q1;q2" — последние прогоны тем
    python -m app.ml.blind_eval score data/bench/blind_ivan.xlsx [data/bench/blind_misha.xlsx]
Результат: data/bench/blind_sheet.xlsx, data/bench/blind_key.json, data/model/blind_eval.json
"""
import argparse
import asyncio
import json
import random
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from app.config import get_settings
from app.db import SessionLocal, init_db
from app.models import Candidate, Document, Run

BENCH = lambda: Path(get_settings().model_dir).parent / "bench"  # noqa: E731
COLS = ["Код", "Запрос", "Технология", "Описание", "Источники", "Релевантность 0–2", "Технология 0–2", "Ранняя стадия 0–2",
        "Русское название 0–2", "Категория", "Комментарий"]
CATEGORIES = "ok,заголовок статьи,продукт,слишком общее,не по теме,зрелое,стандарт,не технология,дубль"


def latest_runs(queries: list[str]) -> list[int]:
    with SessionLocal() as s:
        out = []
        for q in queries:
            r = next((r for r in s.query(Run).filter(Run.status == "done").order_by(Run.id.desc()).limit(400)
                      if r.query.strip().lower() == q.strip().lower()), None)
            if r:
                out.append(r.id)
        return out


def _summary(c: Candidate) -> str:
    d = c.dossier or {}
    if d.get("generated_by") in (None, "template"):
        return ""
    s = re.sub(r"\s*\[\d+(?:[,\s]+\d+)*\]", "", d.get("description", "")).strip()
    return s[:300]


def export(run_ids: list[int], seed: int = 7) -> tuple[Path, Path]:
    rows, key = [], {}
    with SessionLocal() as s:
        for rid in run_ids:
            run = s.get(Run, rid)
            variant = "только LLM" if (run.options or {}).get("seed_candidates") else "конвейер"
            for c in sorted(s.query(Candidate).filter_by(run_id=rid, status="signal"), key=lambda c: c.rank or 99)[:15]:
                docs = [s.get(Document, i) for i in (c.document_ids or [])[:3]]
                rows.append({"query": run.query, "name": c.name_ru or c.name, "en": c.name, "desc": _summary(c),
                             "links": "\n".join(d.url for d in docs if d), "run_id": rid, "cand_id": c.id, "variant": variant, "rank": c.rank})
    random.Random(seed).shuffle(rows)
    wb = Workbook()
    ws = wb.active
    ws.title = "Разметка"
    ws.append(COLS)
    for cell in ws[1]:
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="5B2FB3")
    for i, r in enumerate(rows, 1):
        code = f"K{i:03d}"
        key[code] = {k: r[k] for k in ("run_id", "cand_id", "variant", "rank", "query", "en")}
        name = r["name"] if r["name"] == r["en"] else f"{r['name']} ({r['en']})"
        ws.append([code, r["query"], name, r["desc"], r["links"], None, None, None, None, None, None])
    score = DataValidation(type="whole", operator="between", formula1="0", formula2="2", allow_blank=True)
    cat = DataValidation(type="list", formula1=f'"{CATEGORIES}"', allow_blank=True)
    ws.add_data_validation(score)
    ws.add_data_validation(cat)
    score.add(f"F2:I{len(rows) + 1}")
    cat.add(f"J2:J{len(rows) + 1}")
    for col, w in zip("ABCDEFGHIJK", [7, 30, 40, 60, 50, 12, 12, 12, 12, 18, 30], strict=True):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    rules = wb.create_sheet("Как размечать")
    for line in ["Оценивайте каждую строку независимо. Оценки модели, место и происхождение карточки скрыты намеренно.",
                 "Релевантность: 0 — не по теме запроса, 1 — соседняя тема, 2 — прямо по запросу.",
                 "Технология: 0 — заголовок статьи, организация, продукт или общее слово; 1 — область; 2 — конкретная технология.",
                 "Ранняя стадия: 0 — зрелое, массовое или стандарт; 1 — уже на рынке, но растёт; 2 — исследования, пилоты, первые продукты.",
                 "Русское название: 0 — неверно или калька, 1 — понятно, но коряво, 2 — как в отраслевой статье.",
                 "Категория — главная проблема строки (или ok). Комментарий — одна фраза по желанию."]:
        rules.append([line])
    rules.column_dimensions["A"].width = 140
    out = BENCH()
    out.mkdir(parents=True, exist_ok=True)
    sheet, keyf = out / "blind_sheet.xlsx", out / "blind_key.json"
    wb.save(sheet)
    keyf.write_text(json.dumps(key, ensure_ascii=False, indent=1), encoding="utf-8")
    return sheet, keyf


def _read(path: Path) -> dict[str, dict]:
    ws = load_workbook(path, data_only=True)["Разметка"]
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        vals = [row[i] if isinstance(row[i], int | float) else None for i in (5, 6, 7, 8)]
        if all(v is None for v in vals):
            continue
        out[row[0]] = {"rel": vals[0], "tech": vals[1], "early": vals[2], "name": vals[3], "cat": row[9]}
    return out


def useful(r: dict) -> bool:
    return (r["rel"] or 0) >= 1 and (r["tech"] or 0) == 2 and (r["early"] or 0) >= 1


def kappa(a: list[bool], b: list[bool]) -> float | None:
    n = len(a)
    if not n:
        return None
    po = sum(x == y for x, y in zip(a, b, strict=True)) / n
    pa, pb = sum(a) / n, sum(b) / n
    pe = pa * pb + (1 - pa) * (1 - pb)
    return round((po - pe) / (1 - pe), 3) if pe < 1 else None


def score(paths: list[Path]) -> dict:
    key = json.loads((BENCH() / "blind_key.json").read_text(encoding="utf-8"))
    raters = [_read(p) for p in paths]
    main = raters[0]
    res: dict = {"raters": len(raters), "labelled": len(main), "by_variant": {}, "by_query": {}}
    for code, r in main.items():
        k = key.get(code)
        if not k:
            continue
        v = res["by_variant"].setdefault(k["variant"], {"cards": 0, "useful": 0, "queries": set()})
        v["cards"] += 1
        v["useful"] += useful(r)
        v["queries"].add(k["query"])
        q = res["by_query"].setdefault(f"{k['query']} · {k['variant']}", {"cards": 0, "useful": 0, "problems": {}})
        q["cards"] += 1
        q["useful"] += useful(r)
        if r.get("cat") and r["cat"] != "ok":
            q["problems"][r["cat"]] = q["problems"].get(r["cat"], 0) + 1
    for v in res["by_variant"].values():
        nq = len(v.pop("queries"))
        v["U@15"] = round(v["useful"] / nq, 2) if nq else None
        v["P@15"] = round(v["useful"] / (15 * nq), 3) if nq else None
        v["useful_share"] = round(v["useful"] / v["cards"], 3) if v["cards"] else None
    if len(raters) > 1:
        common = sorted(set(raters[0]) & set(raters[1]))
        res["kappa_useful"] = kappa([useful(raters[0][c]) for c in common], [useful(raters[1][c]) for c in common])
        res["kappa_n"] = len(common)
    (Path(get_settings().model_dir) / "blind_eval.json").write_text(json.dumps(res, ensure_ascii=False, indent=1), encoding="utf-8")
    return res


async def run_holdout(with_llm: bool = False, only: list[str] | None = None) -> list[int]:
    """Живые прогоны тем HOLD-OUT; with_llm — для каждой темы ещё и вариант «только LLM» (для сравнения в одном листе).
    Прогон с исчерпанной квотой OpenAlex (degraded) в лист не берётся: его надо повторить после сброса квоты (03:00 МСК)."""
    from app.pipeline import runner

    qs = only or json.loads((BENCH() / "queries.json").read_text(encoding="utf-8"))["holdout"]
    ids = []
    for q in qs:
        for opts in ([{"bench": "holdout"}] + ([{"bench": "holdout", "seed_candidates": True}] if with_llm else [])):
            with SessionLocal() as s:
                r = Run(query=q, options=opts)
                s.add(r)
                s.commit()
                rid = r.id
            await runner.execute(rid)
            with SessionLocal() as s:
                bad = bool((s.get(Run, rid).stats or {}).get("degraded"))
            print("готово:", q, rid, "(НЕДЕЙСТВИТЕЛЕН: OpenAlex недоступен)" if bad else "", flush=True)
            if not bad:
                ids.append(rid)
    return ids


def main() -> None:
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    e = sub.add_parser("export")
    e.add_argument("--runs", default="")
    e.add_argument("--queries", default="")
    sc = sub.add_parser("score")
    sc.add_argument("files", nargs="+")
    rh = sub.add_parser("run-holdout")
    rh.add_argument("--with-llm", action="store_true", help="добавить вариант «только LLM» по каждой теме")
    rh.add_argument("--only", default="", help="темы через «;» (по умолчанию — весь HOLD-OUT)")
    args = ap.parse_args()
    init_db()
    if args.cmd == "run-holdout":
        ids = asyncio.run(run_holdout(args.with_llm, [q for q in args.only.split(";") if q.strip()] or None))
        print("Лист:", export(ids))
    elif args.cmd == "export":
        ids = [int(x) for x in args.runs.split(",") if x.strip()] or latest_runs([q for q in args.queries.split(";") if q.strip()])
        print("Лист и ключ:", export(ids))
    else:
        print(json.dumps(score([Path(f) for f in args.files]), ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
