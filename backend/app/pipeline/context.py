"""«Под вашу компанию»: какие найденные сигналы связаны с документом компании (стратегией, ТЗ, планом НИОКР) и чем именно.

Без LLM и без GPU: документ режется на абзацы, сигнал описывается своим названием (EN и RU), справкой и заголовками
подтверждающих источников; сходство — TF-IDF по символьным n-граммам (работает для русского и английского текста
и для смешанных терминов). Для каждого связанного сигнала показывается самый близкий абзац документа и общие слова —
связь проверяема глазами. Документ нигде не сохраняется: разбирается в памяти и отбрасывается после ответа.
"""
import base64
import io
import re
import zipfile

import numpy as np
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

MAX_BYTES = 5 * 1024 * 1024
MIN_SCORE = 0.08  # ниже — совпадение случайных буквосочетаний (подобрано на синтетических парах «стратегия ↔ сигналы»)


class ContextError(ValueError):
    pass


def extract_text(filename: str, data: bytes) -> str:
    """Текст из .txt / .md / .docx / .xlsx / .pdf (PDF — если установлен pypdf)."""
    if len(data) > MAX_BYTES:
        raise ContextError("Файл больше 5 МБ")
    name = (filename or "").lower()
    if name.endswith((".txt", ".md", ".csv")):
        for enc in ("utf-8-sig", "cp1251"):
            try:
                return data.decode(enc)
            except UnicodeDecodeError:
                continue
        raise ContextError("Не удалось прочитать текст (кодировка)")
    if name.endswith(".docx"):
        try:
            xml = zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml").decode("utf-8", "ignore")
        except (zipfile.BadZipFile, KeyError) as exc:
            raise ContextError("Файл .docx повреждён") from exc
        xml = re.sub(r"</w:p>", "\n", xml)
        return re.sub(r"<[^>]+>", "", xml)
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        return "\n".join(" ".join(str(v) for v in row if v not in (None, "")) for ws in wb.worksheets for row in ws.iter_rows(values_only=True))
    if name.endswith(".pdf"):
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise ContextError("Для PDF нужен пакет pypdf; сохраните документ как .docx или .txt") from exc
        return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
    raise ContextError("Поддерживаются .txt, .md, .docx, .xlsx, .pdf")


def decode(payload: dict) -> tuple[str, str]:
    """{filename, content_base64} или {filename, text} → (имя, текст)."""
    fn = str(payload.get("filename") or "документ.txt")
    if payload.get("text"):
        text = str(payload["text"])
    else:
        try:
            raw = base64.b64decode(payload.get("content_base64") or "", validate=True)
        except (ValueError, TypeError) as exc:
            raise ContextError("Файл не прочитан") from exc
        text = extract_text(fn, raw)
    text = re.sub(r"[ \t\r\f\v]+", " ", text).strip()
    if len(text) < 40:
        raise ContextError("В документе почти нет текста")
    return fn, text[:400_000]


def paragraphs(text: str) -> list[str]:
    parts = [p.strip() for p in re.split(r"\n\s*\n|\n(?=[-•*\d]+[.)]?\s)|\n", text) if len(p.strip()) >= 40]
    # Длинные абзацы режем по предложениям (до ~600 знаков): цитата должна быть короткой.
    out = []
    for p in parts:
        while len(p) > 600:
            cut = p.rfind(". ", 0, 600)
            cut = cut + 1 if cut > 200 else 600
            out.append(p[:cut].strip())
            p = p[cut:].strip()
        if len(p) >= 40:
            out.append(p)
    return out[:3000]


_STOP = set("это для при как что или над под без его она они так также the and for with from that this are was were into using".split())


def _words(s: str) -> set[str]:
    return {w[:6] for w in re.findall(r"[a-zа-яё]{4,}", s.lower()) if w not in _STOP}


def match(text: str, signals: list[dict], top: int = 10) -> list[dict]:
    """signals: [{id, name, name_ru, profile}] → связанные сигналы по убыванию сходства, с цитатой из документа."""
    paras = paragraphs(text)
    if not paras or not signals:
        return []
    vec = TfidfVectorizer(analyzer="char_wb", ngram_range=(3, 5), sublinear_tf=True, min_df=1)
    vec.fit(paras + [s["profile"] for s in signals])
    P, S = vec.transform(paras), vec.transform([s["profile"] for s in signals])
    sim = cosine_similarity(S, P)  # сигналы × абзацы
    out = []
    for i, s in enumerate(signals):
        j = int(np.argmax(sim[i]))
        score = float(sim[i, j])
        if score < MIN_SCORE:
            continue
        common = sorted(_words(paras[j]) & _words(s["profile"]))[:8]
        out.append({"id": s["id"], "name": s["name"], "name_ru": s.get("name_ru"), "score": round(score, 3),
                    "quote": paras[j][:400], "paragraph": j + 1, "common": common,
                    "strength": "сильная" if score >= 0.2 else "заметная" if score >= 0.12 else "слабая"})
    return sorted(out, key=lambda r: -r["score"])[:top]


def signal_profile(name: str, name_ru: str | None, dossier: dict | None, titles: list[str]) -> str:
    d = dossier or {}
    return " ".join([name, name_ru or "", d.get("description", ""), d.get("application", ""), d.get("advantage", ""), *titles[:6]])
