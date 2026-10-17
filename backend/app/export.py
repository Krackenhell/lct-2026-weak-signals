"""Экспорт результата запроса для аналитика и техкомитета: Excel (в колонках датасета организатора), CSV, JSON, Markdown.

Всё строится только из сохранённого прогона — экспорт воспроизводит то, что показано в интерфейсе.
"""
import csv
import io
import json
import math

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.pipeline.passport import prob_text

# Колонки датасета организатора (id, name, domain, companies, why, stage, trend, score, sources) + наша проверяемость.
COLUMNS = [
    ("rank", "№"), ("name", "Технология"), ("name_en", "Technology (EN)"), ("domain", "Направление (запрос)"),
    ("companies", "Организации из источников"), ("why", "Почему это слабый сигнал"), ("stage", "Стадия (самое твёрдое событие)"),
    ("trend", "Тренд публикаций"), ("score", "Оценка 1–10"), ("probability", "Вероятность по модели"),
    ("passport", "Паспорт доказательств"), ("robustness", "Запас прочности"), ("window", "Окно для России"),
    ("origin", "Происхождение"), ("sources", "Источники"),
]


def _trend(f: dict) -> str:
    works = int(round(math.expm1(f.get("log_total_works", 0) or 0)))
    g = f.get("growth", 0) or 0
    if not works:
        return "Публикаций по точной фразе почти нет"
    word = "растёт быстро" if g > 0.7 else "растёт" if g > 0.2 else "стабильно" if g > -0.2 else "снижается"
    return f"{word}: ~{works} работ, {f.get('recent_share', 0):.0%} — за последние 3 года"


def candidate_row(c, docs_by_id: dict, query: str) -> dict:
    extra = c.contributions or {}
    pp, rb, w = extra.get("passport") or {}, extra.get("robustness") or {}, extra.get("window") or {}
    d = c.dossier or {}
    docs = [docs_by_id[i] for i in (c.document_ids or []) if i in docs_by_id]
    orgs = sorted({doc.source_name for doc in docs if doc.source_name and doc.source_type in ("press", "repository", "news")})[:5]
    return {
        "rank": c.rank or "",
        "name": c.name_ru or c.name,
        "name_en": c.name,
        "domain": query,
        "companies": ", ".join(orgs),
        "why": d.get("why_weak_signal") or c.exclusion_reason or "",
        "stage": pp.get("event", ""),
        "trend": _trend(c.features or {}),
        "score": round((c.probability or 0) * 10) if c.probability is not None else "",
        "probability": prob_text(c.probability),
        "passport": f"{pp['level']}: подтверждено {pp['confirmed']} из {pp['checked']}, источников {pp['domains']}" if pp else "",
        "robustness": rb.get("text", ""),
        "window": f"{w.get('label', '')}. {w.get('text', '')}".strip(". ") if w else "",
        "origin": "из базы отслеживания" if c.origin == "watchlist" else "новый кандидат",
        "sources": "\n".join(f"{doc.title} — {doc.url}" for doc in docs[:8]),
        "status": c.status,
        "exclusion": c.exclusion_reason or "",
    }


def build(run, cands, docs) -> dict:
    docs_by_id = {d.id: d for d in docs}
    by_status = {"signal": [], "watch": [], "excluded": []}
    for c in sorted(cands, key=lambda c: (c.rank or 999, -(c.probability or 0))):
        by_status.setdefault(c.status, []).append(candidate_row(c, docs_by_id, run.query))
    used = {i for c in cands if c.status == "signal" for i in (c.document_ids or [])}
    sources = [
        {"title": d.title, "url": d.url, "date": d.published or "", "type": d.source_type, "language": d.language, "trust": d.trust_level,
         "summary_ru": d.summary_ru or "", "summary_generated": bool(d.summary_generated)}
        for d in docs if d.id in used
    ]
    meta = {
        "query": run.query, "run_id": run.id, "status": run.status,
        "date": run.created_at.isoformat() if run.created_at else None,
        "mode": "мониторинг (с базой отслеживания)" if (run.options or {}).get("use_watchlist") else "чистое обнаружение",
        "providers": run.providers or {}, "stats": {k: v for k, v in (run.stats or {}).items() if k != "http"},
    }
    return {"meta": meta, "signals": by_status["signal"], "watch": by_status["watch"], "excluded": by_status["excluded"], "sources": sources}


def to_json(data: dict) -> bytes:
    return json.dumps(data, ensure_ascii=False, indent=1).encode("utf-8")


def to_csv(data: dict) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow([title for _, title in COLUMNS])
    for row in data["signals"]:
        w.writerow([row[k] for k, _ in COLUMNS])
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM — чтобы Excel открыл кириллицу без настроек


def to_markdown(data: dict) -> bytes:
    m = data["meta"]
    lines = [f"# Слабые сигналы: «{m['query']}»", "", f"Прогон №{m['run_id']} от {(m['date'] or '')[:10]} · режим: {m['mode']} · "
             f"модели: {m['providers'].get('decision_model', '—')}, {m['providers'].get('llm', '—')}", ""]
    if m["status"] == "rejected":
        lines += ["Запрос отклонён: он не похож на технологическое направление.", ""]
    for r in data["signals"]:
        lines += [f"## {r['rank']}. {r['name']}", f"*{r['name_en']}* · вероятность {r['probability']} · {r['passport']}", ""]
        for label, key in (("Почему слабый сигнал", "why"), ("Стадия", "stage"), ("Тренд", "trend"), ("Запас прочности", "robustness"), ("Окно для России", "window")):
            if r[key]:
                lines.append(f"- **{label}:** {r[key]}")
        if r["sources"]:
            lines += ["- **Источники:**"] + [f"  - {s}" for s in r["sources"].splitlines()]
        lines.append("")
    if data["watch"]:
        lines += ["## Наблюдать (мало подтверждённых документов)", ""] + [f"- {r['name']} — {r['passport']}" for r in data["watch"]] + [""]
    return "\n".join(lines).encode("utf-8")


def to_xlsx(data: dict) -> bytes:
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="5B2FB3")

    def sheet(ws, header, rows, widths):
        ws.append(header)
        for cell in ws[1]:
            cell.font, cell.fill = head, fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for r in rows:
            ws.append(r)
        for i, wdt in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = wdt
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "Сигналы"
    widths = [5, 34, 30, 22, 26, 60, 22, 30, 9, 14, 30, 40, 40, 18, 70]
    sheet(ws, [t for _, t in COLUMNS], [[r[k] for k, _ in COLUMNS] for r in data["signals"]], widths)
    sheet(wb.create_sheet("Наблюдать"), ["Технология", "Technology (EN)", "Вероятность по модели", "Паспорт доказательств", "Почему не в топе"],
          [[r["name"], r["name_en"], r["probability"], r["passport"], r["exclusion"]] for r in data["watch"]], [34, 30, 14, 40, 70])
    sheet(wb.create_sheet("Исключённые"), ["Технология", "Technology (EN)", "Вероятность по модели", "Причина исключения"],
          [[r["name"], r["name_en"], r["probability"], r["exclusion"]] for r in data["excluded"]], [34, 30, 14, 80])
    sheet(wb.create_sheet("Источники"), ["Название", "Ссылка", "Дата", "Тип", "Язык", "Доверие", "Резюме на русском", "Резюме сгенерировано LLM"],
          [[s["title"], s["url"], s["date"], s["type"], s["language"], s["trust"], s["summary_ru"], "да" if s["summary_generated"] else ""] for s in data["sources"]],
          [50, 45, 12, 12, 8, 10, 60, 12])
    m = data["meta"]
    about = [["Запрос", m["query"]], ["Прогон", m["run_id"]], ["Дата", m["date"]], ["Режим", m["mode"]],
             ["Decision-модель", m["providers"].get("decision_model", "")], ["LLM", m["providers"].get("llm", "")],
             ["Классификатор", m["providers"].get("classifier", "")]]
    about += [[f"Статистика: {k}", json.dumps(v, ensure_ascii=False) if isinstance(v, dict | list) else v] for k, v in m["stats"].items()]
    sheet(wb.create_sheet("О прогоне"), ["Параметр", "Значение"], about, [30, 90])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


FORMATS = {
    "xlsx": (to_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    "csv": (to_csv, "text/csv; charset=utf-8"),
    "json": (to_json, "application/json"),
    "md": (to_markdown, "text/markdown; charset=utf-8"),
}
