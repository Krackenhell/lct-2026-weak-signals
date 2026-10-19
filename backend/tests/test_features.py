"""Восемь фишек к финалу: лестница «от науки до рынка», «под вашу компанию», подписка, прогноз с проверкой,
доводы «против», карта игроков, выдача по мере подтверждения, слепой замер. Без сети: Brave и OpenAlex подменены."""
import asyncio
import base64
import io
import json
import zipfile
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace

import httpx
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import main as main_mod
from app.collectors import brave
from app.collectors.base import RawDoc
from app.config import get_settings
from app.db import SessionLocal
from app.models import Candidate, Forecast, Run, Subscription, SubscriptionEvent
from app.pipeline import context as ctx_mod
from app.pipeline import forecasts, market, runner, watch
from tests.fakes import FakeWeb


def run(coro):
    return asyncio.run(coro)


# ---------------------------------------------------------------- Brave: разбор ответа, события, бюджет

def _brave_client(results: list[dict], hits: list):
    from app.collectors.http import HttpClient

    def handler(req: httpx.Request):
        hits.append(str(req.url))
        assert req.headers.get("X-Subscription-Token")  # ключ — в заголовке, не в адресе
        return httpx.Response(200, json={"web": {"results": results}})

    h = HttpClient()
    h._client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    h._interval = {}
    import tempfile
    from pathlib import Path

    h._cache_dir = Path(tempfile.mkdtemp())
    return h


def test_brave_search_events_and_budget(monkeypatch):
    monkeypatch.setenv("BRAVE_API_KEY", "test-not-a-key")
    get_settings.cache_clear()
    try:
        hits = []
        h = _brave_client([
            {"title": "Acme raises $20M Series A for solid-state battery pilot line", "url": "https://news.example.com/a", "description": "solid-state battery",
             "page_age": "2026-05-01T00:00:00"},
            {"title": "Solid-state battery: a review", "url": "https://blog.example.com/b", "description": "overview of solid-state battery"},
        ], hits)
        b = brave.Budget(limit=1)
        docs = run(brave.events(h, "solid-state battery", b))
        assert [d.extra["market"] for d in docs] == ["funding"] and docs[0].published == "2026-05-01"
        assert run(brave.search(h, '"solid-state battery" pilot OR contract OR funding OR deployment', 10, budget=b))  # из кэша — без бюджета
        try:
            run(brave.patents(h, "solid-state battery", b))
            raise AssertionError("бюджет должен быть исчерпан")
        except Exception as exc:  # noqa: BLE001
            assert "бюджет" in str(exc)
        assert len(hits) == 1
        assert "test-not-a-key" not in " ".join(hits)
        run(h.close())
    finally:
        get_settings.cache_clear()


def test_market_event_types():
    assert brave.market_event("Startup raises $5M seed round") == "funding"
    assert brave.market_event("Company awarded contract for pilot deployment") == "contract"
    assert brave.market_event("Field trial of sensor completed") == "pilot"
    assert brave.market_event("Series production starts") == "launch"
    assert brave.market_event("Нефтяная компания объявила тендер на закупку") == "contract"


# ---------------------------------------------------------------- лестница и доводы «против»

def test_ladder_orders_steps_from_science_to_market():
    items = [{"type": "preprint", "title": "Paper", "date": "2023-02-01"}, {"type": "grant", "event": "grant", "title": "NSF grant", "date": "2024-01-01"},
             {"type": "news", "event": "funding", "title": "Acme raises seed round", "date": "2025-03-01"},
             {"type": "news", "event": "launch", "title": "Pilot plant opened", "snippet": "field trial", "date": "2026-01-01"}]
    lad = market.ladder(items, epo_patents=2)
    by = {s["key"]: s["n"] for s in lad["steps"]}
    assert by["science"] == 1 and by["patent"] == 2 and by["grant"] == 1 and by["funding"] == 1 and by["pilot"] == 1
    assert lad["reached"] == "pilot" and "пилот" in lad["text"] and "2026" in lad["text"]


class _Jev:
    provider, model = "jev", "jev-1.13.0"

    async def ask(self, state, questions):
        ans = {}
        for k in questions:
            if k.startswith("doc_"):
                ans[k] = {"value": 0.1 if "Unrelated" in state.split(f"[{k[4:]}]")[1].split("\n")[0] else 0.9}
            if k.startswith("against_"):
                ans[k] = {"value": 0.8}
        return SimpleNamespace(answers=ans, ok=True)


def test_market_layer_confirms_web_docs_and_against(monkeypatch):
    monkeypatch.setattr(brave, "enabled", lambda: True)

    async def ev(http, p, b):
        return [RawDoc(url="https://n.example/1", title=f"{p} pilot deployed at refinery", snippet="field trial", published="2026-02-01",
                       source_name="n", source_type="news", collector="brave", extra={"event": "launch", "market": "pilot"}),
                RawDoc(url="https://n.example/2", title="Unrelated contract", snippet="", source_name="n", source_type="news", collector="brave",
                       extra={"event": "launch", "market": "contract"})]

    async def none(*a, **k):
        return []

    async def ag(http, p, b):
        return [RawDoc(url="https://n.example/3", title=f"{p} project discontinued", snippet="", source_name="n", source_type="news", collector="brave", extra={})]

    monkeypatch.setattr(brave, "events", ev)
    monkeypatch.setattr(brave, "patents", none)
    monkeypatch.setattr(brave, "procurement", none)
    monkeypatch.setattr(brave, "against", ag)
    m = run(market.market_layer(None, _Jev(), "smart gas lift valves", "интеллектуальные клапаны газлифта", "gas lift", [], {}, 0, 0, brave.Budget(10)))
    assert m["ladder"]["reached"] == "pilot" and len(m["web_docs"]) == 1  # «Unrelated» Jev не подтвердила
    assert m["against"] and m["against"][0]["why"] == "провал или отказ"


def test_players_from_openalex_and_items(monkeypatch):
    from app.collectors import openalex

    async def inst(http, phrase, year, country=None, n=8):
        return [{"name": "Российская академия наук", "works": 7}] if country else [{"name": "MIT", "works": 30}]

    monkeypatch.setattr(openalex, "institution_counts", inst)
    pl = run(market.players(None, "rhombohedral graphene", 2026, [{"type": "startup", "title": "GrapheneCo: graphene chips", "url": "u"},
                                                                    {"type": "news", "event": "funding", "title": "Acme Labs raises $10M"}]))
    assert pl["institutions"][0]["name"] == "MIT" and pl["russia"][0]["works"] == 7
    assert {c["name"] for c in pl["companies"]} == {"GrapheneCo", "Acme Labs"}


# ---------------------------------------------------------------- «под вашу компанию»

def _docx(text: str) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("word/document.xml", "<w:document><w:body>" + "".join(f"<w:p><w:r><w:t>{p}</w:t></w:r></w:p>" for p in text.split("\n")) + "</w:body></w:document>")
    return buf.getvalue()


def test_company_context_matches_paragraphs():
    text = ("Стратегия НИОКР на 2027 год.\nПлан: повышение добычи на газоконденсатных скважинах за счёт интеллектуальных клапанов газлифта "
            "и мониторинга притока.\nВторое направление — внедрение электронного документооборота в бухгалтерии компании.")
    fn, t = ctx_mod.decode({"filename": "стратегия.docx", "content_base64": base64.b64encode(_docx(text)).decode()})
    sigs = [{"id": 1, "name": "smart gas lift valves", "name_ru": "интеллектуальные клапаны газлифта",
             "profile": ctx_mod.signal_profile("smart gas lift valves", "интеллектуальные клапаны газлифта", {}, ["Smart gas lift valve field trial"])},
            {"id": 2, "name": "rhombohedral graphene", "name_ru": "ромбоэдрический графен",
             "profile": ctx_mod.signal_profile("rhombohedral graphene", "ромбоэдрический графен", {}, ["Superconductivity in graphene"])}]
    res = ctx_mod.match(t, sigs)
    assert res and res[0]["id"] == 1 and "газлифта" in res[0]["quote"] and all(r["id"] != 2 or r["score"] < res[0]["score"] for r in res)


def test_company_context_api_does_not_store_document():
    with SessionLocal() as s:
        r = Run(query="газлифт контекст", status="done")
        s.add(r)
        s.flush()
        s.add(Candidate(run_id=r.id, name="smart gas lift valves", name_ru="интеллектуальные клапаны газлифта", status="signal", rank=1, probability=0.8,
                        relevance=0.9, features={}, contributions={}, answers={}, dossier={}, document_ids=[]))
        s.commit()
        rid = r.id
    with TestClient(main_mod.app) as c:
        res = c.post(f"/api/runs/{rid}/context", json={"filename": "план.txt",
                                                        "text": "Задача на год: опытная эксплуатация интеллектуальных клапанов газлифта на трёх скважинах."}).json()
        assert res["matches"] and res["matches"][0]["name"] == "smart gas lift valves" and "не сохранён" in res["note"]
        assert c.post(f"/api/runs/{rid}/context", json={"filename": "x.exe", "content_base64": base64.b64encode(b"MZ").decode()}).status_code == 400


# ---------------------------------------------------------------- подписка

def _signal(run_id, name, level, reached):
    return Candidate(run_id=run_id, name=name, status="signal", rank=1, probability=0.8, relevance=0.9, features={}, answers={}, dossier={},
                     document_ids=[], contributions={"passport": {"level": level}, "market": {"ladder": {"reached": reached}}})


def test_subscription_compare_reports_strengthened_signals():
    with SessionLocal() as s:
        a, b = Run(query="подписка тест", status="done"), Run(query="подписка тест", status="done")
        s.add_all([a, b])
        s.flush()
        s.add_all([_signal(a.id, "tech one", "B", "science"), _signal(a.id, "tech gone", "B", None),
                   _signal(b.id, "tech one", "A", "funding"), _signal(b.id, "tech new", "B", None)])
        s.commit()
        diff = watch.compare(s, a.id, b.id)
    assert [x["name"] for x in diff["new"]] == ["tech new"] and [x["name"] for x in diff["dropped"]] == ["tech gone"]
    assert diff["stronger"][0]["why"] == ["новая ступень: раунд / инвестиции", "доказательства B → A"]
    assert "Усилились" in watch.summarize(diff)


def test_subscription_refresh_and_api(monkeypatch):
    async def fake_execute(run_id):
        with SessionLocal() as s:
            r = s.get(Run, run_id)
            r.status = "done"
            s.add(_signal(run_id, "tech one", "A", "pilot"))
            s.commit()

    monkeypatch.setattr(runner, "execute", fake_execute)
    with TestClient(main_mod.app) as c:
        assert c.post("/api/subscriptions", json={"query": "тема подписки", "webhook": "http://127.0.0.1/x"}).status_code == 400  # SSRF
        sub = c.post("/api/subscriptions", json={"query": "тема подписки", "every_days": 7}).json()
        ev = run(watch.refresh(sub["id"]))
        assert ev and "Первый расчёт" in ev.summary
        lst = c.get("/api/subscriptions").json()
        assert any(x["id"] == sub["id"] and x["events"] for x in lst)
        with SessionLocal() as s:
            assert s.get(Subscription, sub["id"]).next_at.replace(tzinfo=timezone.utc) > datetime.now(timezone.utc) + timedelta(days=6)
            assert s.query(SubscriptionEvent).filter_by(subscription_id=sub["id"]).count() == 1
        assert c.delete(f"/api/subscriptions/{sub['id']}").json()["ok"]


# ---------------------------------------------------------------- прогноз с проверкой

def test_forecast_record_and_check(monkeypatch):
    from app.collectors import openalex

    counts = {"base": 20, "later": 40}

    async def fake_count(http, flt):
        return counts["later"] if "to_publication_date:2027" in flt else counts["base"]

    monkeypatch.setattr(openalex, "count", fake_count)
    monkeypatch.setattr(brave, "enabled", lambda: False)
    with SessionLocal() as s:
        r = Run(query="прогноз тест", status="done", options={})
        s.add(r)
        s.flush()
        s.add_all([_signal(r.id, "forecast tech growing", "B", "science"), _signal(r.id, "forecast tech flat", "B", "science")])
        s.commit()
        rid = r.id
    t0 = datetime(2026, 10, 20, tzinfo=timezone.utc)
    assert run(forecasts.record(rid, http=SimpleNamespace(), now=t0)) == 2
    assert run(forecasts.record(rid, http=SimpleNamespace(), now=t0)) == 0  # открытый прогноз на ту же технологию не дублируется
    assert run(forecasts.check_due(now=t0 + timedelta(days=10), http=SimpleNamespace()))["checked"] == 0  # срок не наступил

    async def per_name(http, flt):
        return 40 if ("growing" in flt and "2027" in flt) else 20

    monkeypatch.setattr(openalex, "count", per_name)
    res = run(forecasts.check_due(now=t0 + timedelta(days=181), http=SimpleNamespace()))
    assert res["confirmed"] >= 1 and res["not_confirmed"] >= 1
    with SessionLocal() as s:
        f = s.query(Forecast).filter_by(name="forecast tech growing").one()
        assert f.status == "confirmed" and f.result["growth"] == 2.0
    assert forecasts.summary()["checked"] >= 2


# ---------------------------------------------------------------- выдача по мере подтверждения

def test_preliminary_results_during_run(monkeypatch):
    monkeypatch.setattr(runner, "HttpClient", FakeWeb().client)
    with SessionLocal() as s:
        r = Run(query="edge AI предварительно", options={})
        s.add(r)
        s.commit()
        rid = r.id
    run(runner.execute(rid))
    with SessionLocal() as s:
        r = s.get(Run, rid)
        assert r.status == "done", r.error
        prelim = r.stats.get("preliminary") or []
        assert all(p["p"] <= 0.95 and p["passport"] in ("A", "B") for p in prelim)


# ---------------------------------------------------------------- слепой замер

def test_blind_eval_export_and_score(tmp_path, monkeypatch):
    from app.ml import blind_eval

    monkeypatch.setattr(blind_eval, "BENCH", lambda: tmp_path)
    with SessionLocal() as s:
        a, b = Run(query="слепой замер", status="done", options={}), Run(query="слепой замер", status="done", options={"seed_candidates": True})
        s.add_all([a, b])
        s.flush()
        for i in range(3):
            s.add(_signal(a.id, f"pipeline tech {i}", "B", None))
            s.add(_signal(b.id, f"llm tech {i}", "B", None))
        s.commit()
        ids = [a.id, b.id]
    sheet, key = blind_eval.export(ids)
    wb = load_workbook(sheet)
    ws = wb["Разметка"]
    header = [c.value for c in ws[1]]
    assert "Оценка" not in " ".join(map(str, header)) and ws.max_row == 7  # 6 карточек, без оценок модели
    k = json.loads(key.read_text(encoding="utf-8"))
    for row in ws.iter_rows(min_row=2):
        good = k[row[0].value]["variant"] == "конвейер"
        for cell, v in zip(row[5:9], (2, 2, 2 if good else 0, 2), strict=True):
            cell.value = v
        row[9].value = "ok" if good else "зрелое"
    filled = tmp_path / "rater.xlsx"
    wb.save(filled)
    monkeypatch.setattr(blind_eval, "get_settings", lambda: SimpleNamespace(model_dir=tmp_path / "model"))
    (tmp_path / "model").mkdir()
    res = blind_eval.score([filled, filled])
    assert res["by_variant"]["конвейер"]["useful"] == 3 and res["by_variant"]["только LLM"]["useful"] == 0
    assert res["kappa_useful"] == 1.0
