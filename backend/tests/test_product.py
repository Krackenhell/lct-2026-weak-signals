"""Продуктовые функции: экспорт (Excel/CSV/JSON/Markdown), оценки аналитика и дообучение на них,
метрики Prometheus, доступ по токену, ретроспективная проверка без утечки будущего."""
import asyncio
import io
import json
from datetime import date

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import get_settings
from app.db import SessionLocal
from app.main import app
from app.ml import backtest, train
from app.models import Candidate, Document, Feedback, ModelCall, Run
from tests.fakes import FakeWeb


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


@pytest.fixture(scope="module")
def run_id():
    with SessionLocal() as s:
        run = Run(query="экспорт: квантовые сенсоры", status="done", progress=100, options={},
                  stats={"signals": 1, "watch": 1, "http": {"api.openalex.org": {"ok": 3, "cached": 0, "failed": 1}}},
                  providers={"decision_model": "jev:jev-latest", "llm": "gigachat:GigaChat-2-Max", "classifier": "logistic-regression"})
        s.add(run)
        s.flush()
        docs = [Document(run_id=run.id, url=f"https://arxiv.org/abs/{i}", title=f"Quantum sensing {i}", source_name=n, source_type=t, trust_level="medium",
                         collector="x", summary_ru="Резюме" if i == 0 else None, summary_generated=i == 0)
                for i, (n, t) in enumerate([("arXiv", "preprint"), ("Qnami (пресс-релиз)", "press"), ("GitHub", "repository")])]
        s.add_all(docs)
        s.flush()
        pp = {"level": "B", "label": "Достаточно", "checked": 6, "confirmed": 4, "domains": 3, "domain_list": [], "science": 2, "event_level": 4, "event": "пилот / испытания", "hard_event": True}
        feats = {"log_total_works": 3.4, "growth": 0.9, "recent_share": 0.8, **{k: 0.5 for k in ("specific_tech", "novelty")}}
        sig = Candidate(run_id=run.id, name="quantum sensing", name_ru="Квантовые сенсоры", status="signal", rank=1, probability=0.97, relevance=0.9, features=feats,
                        contributions={"items": [], "methods": ["llm-extraction"], "passport": pp,
                                       "robustness": {"level": "stable", "label": "устойчивый", "text": "Без довода «Рост» вероятность станет 62%.", "p_without": 0.62},
                                       "window": {"status": "open", "label": "Окно открыто", "text": "В мире 30 работ, в России 0."}},
                        dossier={"why_weak_signal": "Мало работ, быстрый рост."}, document_ids=[d.id for d in docs])
        watch = Candidate(run_id=run.id, name="thin tech", status="watch", probability=0.96, exclusion_code="watch", exclusion_reason="мало подтверждений",
                          contributions={"items": [], "passport": {**pp, "level": "C", "confirmed": 1}}, document_ids=[])
        exc = Candidate(run_id=run.id, name="iso 20022", status="excluded", probability=0.2, exclusion_code="standard", exclusion_reason="Отраслевой стандарт", contributions={"items": []})
        s.add_all([sig, watch, exc, ModelCall(run_id=run.id, provider="jev", model="jev-latest", purpose="x", subject="y", latency_ms=120, ok=True)])
        s.commit()
        return run.id


# ---------------------------------------------------------------- экспорт

def test_export_xlsx_in_dataset_columns(client, run_id):
    r = client.get(f"/api/runs/{run_id}/export?format=xlsx")
    assert r.status_code == 200 and f"run-{run_id}.xlsx" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Сигналы", "Наблюдать", "Исключённые", "Источники", "О прогоне"]
    ws = wb["Сигналы"]
    header = [c.value for c in ws[1]]
    assert header[:9] == ["№", "Технология", "Technology (EN)", "Направление (запрос)", "Организации из источников", "Почему это слабый сигнал",
                          "Стадия (самое твёрдое событие)", "Тренд публикаций", "Балл места (не больше 0,95)"]
    row = {h: c.value for h, c in zip(header, ws[2], strict=True)}
    # Шкалы «1–10» больше нет (она давала «10» — ту же ложную точность, что «100%»): балл = min(p; 0,95) × (0,6 + 0,4 × rel).
    assert row["Технология"] == "Квантовые сенсоры" and row["Оценка классификатора"] == "выше 95%"
    assert row["Балл места (не больше 0,95)"] == round(0.95 * (0.6 + 0.4 * 0.9), 2) <= 0.95
    assert "100%" not in json.dumps(row, ensure_ascii=False) and 10 not in row.values()
    assert row["Стадия (самое твёрдое событие)"] == "пилот / испытания" and row["Паспорт доказательств"].startswith("B: подтверждено 4 из 6")
    assert "Qnami" in row["Организации из источников"] and "arXiv" not in row["Организации из источников"]
    assert row["Тренд публикаций"].startswith("растёт быстро")
    assert wb["Наблюдать"]["A2"].value == "thin tech" and wb["Исключённые"]["D2"].value == "Отраслевой стандарт"
    assert wb["Источники"].max_row == 4 and wb["Источники"]["H2"].value == "да"
    about = {r[0].value: r[1].value for r in wb["О прогоне"].iter_rows(min_row=2)}
    assert about["Decision-модель"] == "jev:jev-latest" and "Статистика: http" not in about  # служебная статистика не выгружается


def test_export_csv_json_md(client, run_id):
    csv_ = client.get(f"/api/runs/{run_id}/export?format=csv")
    text = csv_.content.decode("utf-8")
    assert text.startswith("﻿№;Технология") and "Квантовые сенсоры" in text  # BOM для Excel
    data = client.get(f"/api/runs/{run_id}/export?format=json").json()
    assert [r["name"] for r in data["signals"]] == ["Квантовые сенсоры"] and data["meta"]["mode"] == "чистое обнаружение"
    md = client.get(f"/api/runs/{run_id}/export?format=md").content.decode("utf-8")
    assert md.startswith("# Слабые сигналы: «экспорт: квантовые сенсоры»") and "## 1. Квантовые сенсоры" in md and "## Наблюдать" in md


def test_export_errors(client, run_id):
    assert client.get(f"/api/runs/{run_id}/export?format=pdf").status_code == 400
    assert client.get("/api/runs/999999/export").status_code == 404


# ---------------------------------------------------------------- оценки аналитика

def test_feedback_is_saved_shown_and_counted(client, run_id):
    with SessionLocal() as s:
        sig = s.query(Candidate).filter_by(run_id=run_id, status="signal").one()
        exc = s.query(Candidate).filter_by(run_id=run_id, status="excluded").one()
    assert client.post(f"/api/candidates/{sig.id}/feedback", json={"verdict": "noise"}).json()["label"] == "Завышенные ожидания без подтверждений"
    assert client.post(f"/api/candidates/{sig.id}/feedback", json={"verdict": "signal", "comment": "согласен"}).status_code == 200
    assert client.post(f"/api/candidates/{exc.id}/feedback", json={"verdict": "mature"}).status_code == 200
    assert client.post(f"/api/candidates/{sig.id}/feedback", json={"verdict": "maybe"}).status_code == 422
    assert client.post("/api/candidates/999999/feedback", json={"verdict": "signal"}).status_code == 404
    c = client.get(f"/api/candidates/{sig.id}").json()
    assert c["feedback"] == {"verdict": "signal", "label": "Слабый сигнал", "comment": "согласен"}  # последнее решение
    st = client.get("/api/feedback/stats").json()
    assert st["total"] >= 3 and st["by_verdict"]["signal"] >= 1 and 0 <= st["agreement"] <= 1


def test_retrain_uses_feedback_from_same_decision_model(run_id):
    with SessionLocal() as s:
        cand = s.query(Candidate).filter_by(run_id=run_id, status="signal").one()
        s.add(Feedback(candidate_id=cand.id, run_id=run_id, verdict="signal", name="ok", features={"growth": 1.0, "novelty": 0.9}, provider="jev:jev-latest", n_docs=5))
        s.add(Feedback(candidate_id=cand.id + 1, run_id=run_id, verdict="noise", name="rules", features={}, provider="heuristic:rules-v1", n_docs=5))
        s.add(Feedback(candidate_id=cand.id + 2, run_id=run_id, verdict="mature", name="thin", features={}, provider="jev:jev-latest", n_docs=1))
        s.commit()
    rows, info = train.feedback_rows("jev:jev-latest")
    assert info["other_provider"] >= 1 and info["no_evidence"] >= 1  # чужая decision-модель и примеры без доказательств не берутся
    mine = [r for r in rows if r["name"] == "ok"]
    assert mine and mine[0]["label"] == 1 and mine[0]["numeric"]["growth"] == 1.0 and mine[0]["answers"]["novelty"] == 0.9


# ---------------------------------------------------------------- метрики и доступ

def test_prometheus_metrics(client, run_id):
    with SessionLocal() as s:
        cid = s.query(Candidate).filter_by(run_id=run_id, status="signal").one().id
        s.add(Feedback(candidate_id=cid, run_id=run_id, verdict="signal", provider="jev:jev-latest"))
        s.commit()
    text = client.get("/api/metrics").text
    assert "# TYPE weak_signals_runs gauge" in text and 'weak_signals_runs{status="done"}' in text
    assert 'weak_signals_model_calls{provider="jev",ok="true"}' in text
    assert 'weak_signals_source_failure_ratio{host="api.openalex.org"}' in text
    assert 'weak_signals_feedback{verdict="signal"}' in text


def test_token_protects_writes_but_not_reads(client, run_id, monkeypatch):
    monkeypatch.setattr(get_settings(), "api_token", "s3cret")
    assert client.post("/api/search", json={"query": "edge AI"}).status_code == 401
    assert client.post("/api/search", json={"query": "edge AI"}, headers={"X-API-Token": "wrong"}).status_code == 401
    with SessionLocal() as s:
        cid = s.query(Candidate).filter_by(run_id=run_id, status="signal").one().id
    assert client.post(f"/api/candidates/{cid}/feedback", json={"verdict": "signal"}).status_code == 401
    assert client.post(f"/api/candidates/{cid}/feedback", json={"verdict": "signal"}, headers={"X-API-Token": "s3cret"}).status_code == 200
    assert client.get(f"/api/runs/{run_id}").status_code == 200  # чтение результатов открыто
    assert client.get("/api/health").status_code == 200


# ---------------------------------------------------------------- ретроспективная проверка

def test_backtest_cuts_off_the_future():
    seen = {}

    def handler(r):
        u = str(r.url)
        if "openalex" in u:
            return httpx.Response(200, json={"group_by": [{"key": "2019", "count": 5}, {"key": "2020", "count": 20}, {"key": "2024", "count": 5000}]})
        if "arxiv" in u:
            seen["arxiv"] = r.url.params["search_query"]
            return httpx.Response(200, text='<?xml version="1.0"?><feed xmlns="http://www.w3.org/2005/Atom"></feed>')
        if "algolia" in u:
            seen["hn"] = r.url.params["numericFilters"]
            return httpx.Response(200, json={"hits": [], "nbHits": 3})
        if "github" in u:
            seen["gh"] = r.url.params["q"]
            return httpx.Response(200, json={"total_count": 4, "items": []})
        if "wikimedia" in u:
            seen["views"] = u
            return httpx.Response(200, json={"items": []})  # в 2020 статьи ещё не было
        return httpx.Response(200, json={"query": {"search": [{"title": "Vector database"}]}})

    web = FakeWeb()
    h = web.client()
    h._client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    feats, ev = asyncio.run(backtest.as_of(h, "vector database", date(2021, 1, 1)))
    assert ev.yearly == {2019: 5, 2020: 20}  # 2024 год отрезан
    assert seen["arxiv"].endswith("submittedDate:[199101010000 TO 202101010000]")
    assert seen["hn"].endswith(",created_at_i<1609459200") and seen["gh"].endswith("created:2020-01-02..2021-01-01")
    assert "/20200101/20201231" in seen["views"] and ev.wiki["exists"] is False  # статья, написанная позже, не «протекает» в прошлое
    assert feats["log_total_works"] > 0 and feats["wiki_exists"] == 0.0


def test_backtest_metrics_compare_with_baselines():
    res = [{"label": 1, "p": 0.9, "features": {"log_total_works": 1.0, "growth": 1.0}},
           {"label": 1, "p": 0.8, "features": {"log_total_works": 2.0, "growth": 0.1}},
           {"label": 0, "p": 0.2, "features": {"log_total_works": 5.0, "growth": 0.8}},
           {"label": 0, "p": 0.1, "features": {"log_total_works": 4.0, "growth": -0.5}}]
    m = backtest.evaluate(res)
    assert m["n"] == 4 and m["n_trend"] == 2
    assert m["model"]["roc_auc"] == 1.0 and m["model"]["precision_at_2"] == 1.0
    assert m["baseline_more_works"]["roc_auc"] == 0.0  # «больше работ» здесь ошибается полностью
    assert set(m) == {"n", "n_trend", "model", "baseline_more_works", "baseline_growth"}


def test_backtest_cases_are_balanced_and_documented():
    labels = [c[1] for c in backtest.CASES]
    assert len(backtest.CASES) >= 25 and 0.35 <= sum(labels) / len(labels) <= 0.65
    assert all(note for *_, note in backtest.CASES) and len({c[0] for c in backtest.CASES}) == len(backtest.CASES)


def test_backtest_report_is_written(monkeypatch, tmp_path):
    monkeypatch.setattr(get_settings(), "model_dir", tmp_path)
    rows = json.loads((get_settings().seed_dir.parent / "model" / "training_rows.json").read_text(encoding="utf-8"))
    (tmp_path / "training_rows.json").write_text(json.dumps(rows), encoding="utf-8")
    monkeypatch.setattr(backtest, "CASES", backtest.CASES[:2] + backtest.CASES[-2:])
    monkeypatch.setattr(backtest, "HttpClient", FakeWeb().client)
    # Четыре темы — меньше минимума: метрики честно не публикуются.
    rep = asyncio.run(backtest.run(date(2021, 1, 1)))
    saved = json.loads((tmp_path / "backtest.json").read_text(encoding="utf-8"))
    assert saved["cutoff"] == "2021-01-01" and len(saved["results"]) == 4 and "labels_note" in saved
    assert rep["status"] == "incomplete" and rep["metrics"] is None
    monkeypatch.setattr(backtest, "MIN_COMPLETE", 4)
    rep = asyncio.run(backtest.run(date(2021, 1, 1)))
    assert rep["status"] == "ok" and rep["metrics"]["n"] == 4 and all(0 <= r["p"] <= 1 for r in rep["results"])


def test_backtest_without_openalex_publishes_no_metrics(monkeypatch, tmp_path):
    monkeypatch.setattr(get_settings(), "model_dir", tmp_path)
    rows = json.loads((get_settings().seed_dir.parent / "model" / "training_rows.json").read_text(encoding="utf-8"))
    (tmp_path / "training_rows.json").write_text(json.dumps(rows), encoding="utf-8")
    monkeypatch.setattr(backtest, "MIN_COMPLETE", 2)
    from app.collectors import http as http_mod

    async def no_sleep(_):
        return None

    monkeypatch.setattr(http_mod.asyncio, "sleep", no_sleep)
    monkeypatch.setattr(backtest, "HttpClient", FakeWeb(down={"api.openalex.org"}).client)
    monkeypatch.setattr(backtest, "CASES", backtest.CASES[:2] + backtest.CASES[-2:])
    rep = asyncio.run(backtest.run(date(2021, 1, 1)))
    # Без истории публикаций модель видела бы «нет данных» вместо прошлого темы — такой результат не выдаём за проверку.
    assert rep["incomplete"] == 4 and rep["status"] == "incomplete" and rep["metrics"] is None
